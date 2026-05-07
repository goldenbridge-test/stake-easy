"""
Génère QA_Soutenance_GABIAM.docx
Document Q&A pour la soutenance
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "QA_Soutenance_GABIAM.docx")
doc = Document()

for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.5)

NAVY  = RGBColor(0x0A,0x16,0x28)
BLUE  = RGBColor(0x1A,0x56,0x9E)
GOLD  = RGBColor(0xD4,0xA0,0x17)
WHITE = RGBColor(0xFF,0xFF,0xFF)
GREY  = RGBColor(0x55,0x55,0x55)
GREEN = RGBColor(0x0A,0x7A,0x3A)
DARK  = RGBColor(0x1A,0x1A,0x2E)

def cell_bg(cell, hex_c):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto')
    shd.set(qn('w:fill'), hex_c); tcPr.append(shd)

def no_borders(tbl):
    t = tbl._tbl
    pr = t.find(qn('w:tblPr'))
    if pr is None: pr = OxmlElement('w:tblPr'); t.insert(0,pr)
    b = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement(f'w:{n}'); e.set(qn('w:val'),'none'); b.append(e)
    pr.append(b)

def sf(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"; run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic
    if color: run.font.color.rgb = color

def sp(pt=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(pt)

# ── TITRE ─────────────────────────────────────────────────────────────
t_hdr = doc.add_table(rows=1, cols=1)
no_borders(t_hdr)
cell_bg(t_hdr.cell(0,0), "0A1628")
ph = t_hdr.cell(0,0).paragraphs[0]
ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
ph.paragraph_format.space_before = Pt(14)
ph.paragraph_format.space_after  = Pt(14)
sf(ph.add_run("QUESTIONS – RÉPONSES"), 20, True, WHITE)

sp(4)
t_gold = doc.add_table(rows=1, cols=1)
no_borders(t_gold)
cell_bg(t_gold.cell(0,0), "D4A017")
pg = t_gold.cell(0,0).paragraphs[0]
pg.paragraph_format.space_before = Pt(2); pg.paragraph_format.space_after = Pt(2)

sp(6)
p_sub = doc.add_paragraph(); p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sf(p_sub.add_run("Préparation soutenance  •  GABIAM Dédé Honorine  •  StakeEasy / GoldenBridge"), 10, False, GREY, True)
sp(14)

# ── Helper QA ─────────────────────────────────────────────────────────
_num = [0]
def qa(question, reponse, theme=""):
    _num[0] += 1
    # Numéro + Question (fond bleu)
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    cell_bg(t.cell(0,0), "1A569E")
    pq = t.cell(0,0).paragraphs[0]
    pq.paragraph_format.space_before = Pt(8)
    pq.paragraph_format.space_after  = Pt(8)
    pq.paragraph_format.left_indent  = Cm(0.3)
    if theme:
        sf(pq.add_run(f"[{theme}]  "), 9, False, RGBColor(0xD4,0xA0,0x17))
    sf(pq.add_run(f"Q{_num[0]}.  {question}"), 11, True, WHITE)

    # Réponse
    t2 = doc.add_table(rows=1, cols=1)
    no_borders(t2)
    cell_bg(t2.cell(0,0), "F0F4FA")
    pr = t2.cell(0,0).paragraphs[0]
    pr.paragraph_format.space_before = Pt(8)
    pr.paragraph_format.space_after  = Pt(10)
    pr.paragraph_format.left_indent  = Cm(0.5)
    pr.paragraph_format.right_indent = Cm(0.3)
    sf(pr.add_run("↳  "), 11, True, GOLD)
    sf(pr.add_run(reponse), 11, False, DARK)
    sp(10)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 1 — CONNEXION METAMASK
# ═══════════════════════════════════════════════════════════════════════
p_sec = doc.add_paragraph()
sf(p_sec.add_run("◆  CONNEXION WALLET & METAMASK"), 13, True, BLUE)
p_sec.paragraph_format.space_after = Pt(6)

qa("Comment votre application détecte-t-elle MetaMask ?",
   "MetaMask injecte automatiquement un objet 'window.ethereum' dans le navigateur. "
   "Notre hook useWeb3.ts vérifie si cet objet existe au chargement. "
   "S'il n'existe pas, on affiche 'Installez MetaMask'. S'il existe, on peut communiquer avec le wallet.",
   "MetaMask")

qa("Quelle est la différence entre eth_accounts et eth_requestAccounts ?",
   "eth_accounts est silencieux : il retourne les comptes déjà autorisés sans afficher de popup. "
   "On l'utilise au chargement pour auto-connecter si l'utilisateur s'était déjà connecté. "
   "eth_requestAccounts déclenche la popup d'autorisation MetaMask — on l'utilise uniquement quand l'utilisateur clique sur 'Link Wallet'.",
   "MetaMask")

qa("Quel est le rôle de ethers.js dans la connexion ?",
   "ethers.js est une bibliothèque qui sert de couche d'abstraction entre notre code React et MetaMask. "
   "On passe window.ethereum à ethers.providers.Web3Provider, ce qui crée un provider. "
   "Ensuite getSigner() donne accès au compte actif, et getAddress() retourne l'adresse publique. "
   "Sans ethers.js on devrait écrire beaucoup plus de code bas niveau.",
   "ethers.js")

qa("Que se passe-t-il si l'utilisateur change de réseau ou de compte ?",
   "On écoute les événements accountsChanged et chainChanged émis par MetaMask. "
   "Si le compte change, on met à jour le state React. "
   "Si le réseau change, on recharge la page entière via window.location.reload() pour réinitialiser tous les contrats.",
   "MetaMask")

sp(6)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 2 — STAKING
# ═══════════════════════════════════════════════════════════════════════
p_sec2 = doc.add_paragraph()
sf(p_sec2.add_run("◆  STAKING & SMART CONTRACTS"), 13, True, BLUE)
p_sec2.paragraph_format.space_after = Pt(6)

qa("Expliquez le processus de staking de bout en bout.",
   "Il y a 2 étapes. D'abord on-chain : une approbation ERC-20 (approve) qui autorise le contrat TokenFarm "
   "à prendre les tokens du wallet, puis la transaction stakeTokens qui transfère réellement les tokens. "
   "Chaque étape génère une popup MetaMask que l'utilisateur doit signer. "
   "Ensuite hors-chaîne : une fois la blockchain confirmée, on enregistre la position dans Django "
   "avec le hash de transaction, la durée choisie et la date de fin calculée.",
   "Staking")

qa("Pourquoi y a-t-il 2 transactions pour staker et pas une seule ?",
   "C'est une règle de sécurité du standard ERC-20. Un contrat ne peut pas prendre vos tokens "
   "sans votre autorisation explicite. L'approve donne cette autorisation pour un montant précis. "
   "C'est une protection pour l'utilisateur : si le contrat est malveillant, "
   "il ne peut prendre que le montant que vous avez approuvé, pas tout votre wallet.",
   "Staking")

qa("Qu'est-ce qu'un smart contract et comment interagit-on avec lui ?",
   "Un smart contract est un programme déployé sur la blockchain qui s'exécute automatiquement "
   "selon des règles prédéfinies, sans intermédiaire. On interagit via son ABI (Application Binary Interface) "
   "qui décrit les fonctions disponibles. Dans notre code, on crée un objet Contract avec l'adresse "
   "du contrat et son ABI, puis on appelle ses fonctions comme des méthodes JavaScript normales.",
   "Blockchain")

qa("Quelle est la différence entre une fonction de lecture et une transaction ?",
   "Une fonction de lecture (view/pure en Solidity) ne modifie pas la blockchain, "
   "elle est gratuite et instantanée — ex: stakingBalance(), tokenIsAllowed(). "
   "Une transaction modifie l'état de la blockchain, coûte du gas (frais), "
   "et nécessite une signature MetaMask — ex: stakeTokens(), approve().",
   "Blockchain")

qa("Que contient le fichier map.json ?",
   "C'est notre registre d'adresses de contrats par réseau. Pour chaque chainId "
   "(11155111 = Sepolia, 1337 = réseau local), il stocke les adresses des contrats "
   "TokenFarm et GoldenToken déployés. Cela permet de changer de réseau sans modifier le code : "
   "on lit juste l'adresse correspondante au chainId détecté.",
   "Architecture")

qa("Comment fonctionne le sélecteur de durée et les APY ?",
   "On a défini un tableau DURATION_OPTIONS avec 4 options : 1 an à 8%, 2 ans à 10%, "
   "3 ans à 12%, 4 ans à 15%. C'est une donnée frontend — le choix est stocké dans le state "
   "durationYears et envoyé à Django lors de l'enregistrement. "
   "La blockchain elle-même ne connaît pas la durée : c'est notre backend qui calcule "
   "la end_date et gère la logique de durée.",
   "Staking")

sp(6)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 3 — BACKEND DJANGO / JWT
# ═══════════════════════════════════════════════════════════════════════
p_sec3 = doc.add_paragraph()
sf(p_sec3.add_run("◆  API DJANGO & AUTHENTIFICATION JWT"), 13, True, BLUE)
p_sec3.paragraph_format.space_after = Pt(6)

qa("Comment fonctionne l'authentification JWT dans votre application ?",
   "JWT (JSON Web Token) fonctionne en 2 tokens : un access token (courte durée, ~5 min) "
   "et un refresh token (longue durée, ~7 jours). "
   "À la connexion, Django retourne les deux. À chaque appel API, on envoie l'access token "
   "dans le header Authorization. Quand il expire (erreur 401), on utilise le refresh token "
   "pour en obtenir un nouveau automatiquement, sans demander à l'utilisateur de se reconnecter.",
   "Django/JWT")

qa("Qu'est-ce qu'un endpoint API et comment l'avez-vous consommé ?",
   "Un endpoint est une URL qui expose une fonctionnalité backend. Par exemple "
   "POST /api/blockchain/stake-tokens/stake/ pour enregistrer un staking. "
   "On les consomme via fetch() ou axios en envoyant les données en JSON avec le token JWT. "
   "J'ai centralisé tous les appels dans des fichiers de service (blockchainApi.ts, api.ts) "
   "pour ne pas dupliquer la logique d'authentification partout.",
   "Django/JWT")

qa("Vous n'aviez pas de connaissances Django au départ. Comment avez-vous appris ?",
   "J'ai suivi des formations sur Udemy sur Django REST Framework et l'authentification JWT, "
   "avec l'accord de mon tuteur M. Hospice KAKE. J'ai aussi beaucoup appris en lisant "
   "la documentation officielle et en déboguant les erreurs réelles : erreurs 400, 401, 502. "
   "Cette démarche d'auto-formation est ce qui m'a permis de comprendre les serializers, "
   "les vues basées sur les classes, et la logique des permissions.",
   "Django/JWT")

sp(6)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 4 — FEDAPAY
# ═══════════════════════════════════════════════════════════════════════
p_sec4 = doc.add_paragraph()
sf(p_sec4.add_run("◆  PAIEMENT FEDAPAY"), 13, True, BLUE)
p_sec4.paragraph_format.space_after = Pt(6)

qa("Expliquez le flux de paiement FedaPay.",
   "1) L'utilisateur clique 'Commencer' sur un cours payant. "
   "2) Le frontend appelle notre API Django pour créer un checkout FedaPay (montant, info client). "
   "3) Django contacte l'API FedaPay et retourne une checkout_url. "
   "4) On redirige l'utilisateur vers cette URL (page FedaPay : MoMo ou carte en FCFA). "
   "5) Après paiement, FedaPay redirige vers notre page de retour. "
   "6) Cette page poll l'API pour vérifier le statut de la transaction. "
   "7) Si approved, l'enrollment est déclenché automatiquement.",
   "FedaPay")

qa("Pourquoi avoir choisi FedaPay plutôt que Stripe ou PayPal ?",
   "FedaPay est adapté au marché africain : il supporte le mobile money (MTN MoMo, Moov), "
   "les transactions en FCFA, et est intégré dans l'écosystème béninois. "
   "Stripe et PayPal sont peu accessibles en Afrique de l'Ouest sans carte internationale. "
   "FedaPay permet d'inclure des utilisateurs qui n'ont pas de carte bancaire.",
   "FedaPay")

sp(6)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 5 — ARCHITECTURE FRONTEND
# ═══════════════════════════════════════════════════════════════════════
p_sec5 = doc.add_paragraph()
sf(p_sec5.add_run("◆  ARCHITECTURE FRONTEND & REACT"), 13, True, BLUE)
p_sec5.paragraph_format.space_after = Pt(6)

qa("Qu'est-ce qu'un hook React et pourquoi avez-vous utilisé useWeb3 ?",
   "Un hook est une fonction React qui encapsule une logique réutilisable avec accès au state. "
   "useWeb3 centralise toute la logique blockchain : connexion wallet, staking, lecture des balances. "
   "Plutôt que de réécrire ce code dans chaque composant, tous les composants qui ont besoin "
   "du wallet importent simplement useWeb3() et accèdent aux fonctions dont ils ont besoin.",
   "React")

qa("Pourquoi avoir utilisé TypeScript plutôt que JavaScript ?",
   "TypeScript ajoute le typage statique à JavaScript. Cela nous permet de détecter les erreurs "
   "à la compilation plutôt qu'à l'exécution. Dans un contexte blockchain où une mauvaise "
   "adresse ou un mauvais type peut faire échouer une transaction, "
   "le typage est particulièrement précieux pour éviter les bugs silencieux.",
   "React")

qa("Comment gérez-vous l'état de connexion dans toute l'application ?",
   "useWeb3 stocke account, isConnected, provider et chainId dans le state local du hook. "
   "Les composants qui importent useWeb3 récupèrent ces valeurs. "
   "Pour l'authentification Django (login/JWT), on utilise un AuthContext React "
   "qui fournit user, isLoggedIn et logout à tous les composants via le contexte React.",
   "React")

qa("Comment avez-vous géré les conflits Git lors du merge de branches ?",
   "Un conflit Git survient quand 2 développeurs modifient la même partie d'un fichier. "
   "On doit ouvrir les fichiers en conflit, choisir quelle version garder ou fusionner les deux, "
   "supprimer les marqueurs de conflit (<<<, ===, >>>), puis faire git add pour marquer "
   "le fichier comme résolu, et git commit pour finaliser le merge.",
   "Git")

sp(6)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 6 — QUESTIONS GÉNÉRALES
# ═══════════════════════════════════════════════════════════════════════
p_sec6 = doc.add_paragraph()
sf(p_sec6.add_run("◆  QUESTIONS GÉNÉRALES / BILAN"), 13, True, BLUE)
p_sec6.paragraph_format.space_after = Pt(6)

qa("Quelle a été votre principale difficulté technique ?",
   "La synchronisation entre le frontend et les smart contracts. Chaque nouveau déploiement "
   "du contrat Solidity génère une nouvelle adresse qu'il faut mettre à jour dans map.json. "
   "Sans cette mise à jour, ethers.js appelle le mauvais contrat et retourne "
   "'execution reverted: Token currently isn't allowed'. "
   "J'ai appris à diagnostiquer ça en lisant les logs MetaMask et Sepolia Etherscan.",
   "Bilan")

qa("Qu'avez-vous appris que vous ne saviez pas avant ce stage ?",
   "Trois choses principales : premièrement, Django REST Framework et la consommation d'API "
   "avec JWT, que j'ai appris via Udemy. Deuxièmement, ethers.js et l'interaction avec "
   "des smart contracts — la différence entre lecture on-chain et transaction, le concept "
   "d'approve ERC-20. Troisièmement, le travail en équipe à distance sur un vrai projet "
   "de production avec coordination entre frontend, backend et smart contracts.",
   "Bilan")

qa("Qu'est-ce que la DeFi et pourquoi ce projet a-t-il un impact social ?",
   "DeFi signifie Finance Décentralisée : des services financiers (staking, prêts, échanges) "
   "qui fonctionnent sur la blockchain sans banque ni intermédiaire. "
   "En Afrique de l'Ouest, beaucoup de personnes n'ont pas accès aux banques traditionnelles. "
   "StakeEasy leur permet de faire fructifier leur épargne via le staking, "
   "en utilisant simplement un smartphone et du mobile money via FedaPay.",
   "Bilan")

qa("Si vous deviez refaire ce projet, que changeriez-vous ?",
   "J'aurais voulu apprendre Solidity en amont pour mieux comprendre les contraintes du smart contract "
   "dès le départ. J'aurais aussi mis en place des tests automatisés (Jest pour React, "
   "Hardhat pour les contrats) pour éviter les régressions lors des nouvelles fonctionnalités. "
   "Et j'aurais documenté les endpoints API dans un fichier Swagger dès le début "
   "pour améliorer la communication avec le backend.",
   "Bilan")

# ── FOOTER ────────────────────────────────────────────────────────────
doc.add_paragraph()
t_ft = doc.add_table(rows=1, cols=1)
no_borders(t_ft)
cell_bg(t_ft.cell(0,0), "0A1628")
pf = t_ft.cell(0,0).paragraphs[0]
pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
pf.paragraph_format.space_before = Pt(10)
pf.paragraph_format.space_after  = Pt(10)
sf(pf.add_run("GABIAM Dédé Honorine  •  Soutenance 2025-2026  •  StakeEasy / GoldenBridge"), 9, False, WHITE, True)

doc.save(OUT)
print(f"✅ {OUT}")
