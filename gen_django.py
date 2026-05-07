"""
Génère Django_Soutenance_GABIAM.docx
Comprendre Django REST Framework pour la soutenance
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Django_Soutenance_GABIAM.docx")
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.5)

NAVY  = RGBColor(0x0A,0x16,0x28); NAVY_HEX  = "0A1628"
BLUE  = RGBColor(0x1A,0x56,0x9E); BLUE_HEX  = "1A569E"
GOLD  = RGBColor(0xD4,0xA0,0x17); GOLD_HEX  = "D4A017"
GREEN = RGBColor(0x0A,0x6A,0x2A); GREEN_HEX = "0A6A2A"
WHITE = RGBColor(0xFF,0xFF,0xFF)
GREY  = RGBColor(0x55,0x55,0x55)
DARK  = RGBColor(0x1A,0x1A,0x2E)
LGREEN = "E8F5EE"
LBLUE  = "EEF4FC"
LYELL  = "FDF8E8"

def cell_bg(cell, hx):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto')
    shd.set(qn('w:fill'), hx); tcPr.append(shd)

def no_b(tbl):
    t = tbl._tbl
    pr = t.find(qn('w:tblPr'))
    if pr is None: pr = OxmlElement('w:tblPr'); t.insert(0,pr)
    b = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement(f'w:{n}'); e.set(qn('w:val'),'none'); b.append(e)
    pr.append(b)

def sf(run, size=11, bold=False, color=None, italic=False):
    run.font.name="Calibri"; run.font.size=Pt(size)
    run.font.bold=bold; run.font.italic=italic
    if color: run.font.color.rgb=color

def sp(pt=6):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(pt)

def section_title(txt):
    sp(4)
    t = doc.add_table(rows=1, cols=1); no_b(t)
    cell_bg(t.cell(0,0), NAVY_HEX)
    p = t.cell(0,0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(10)
    p.paragraph_format.left_indent  = Cm(0.4)
    sf(p.add_run(txt), 13, True, WHITE)
    sp(6)

def concept(titre, corps, bg=LBLUE, titre_color=None):
    """Bloc concept : titre bleu + corps"""
    t = doc.add_table(rows=1, cols=1); no_b(t)
    cell_bg(t.cell(0,0), BLUE_HEX)
    p = t.cell(0,0).paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Cm(0.4)
    sf(p.add_run(titre), 11, True, WHITE)
    # Corps
    t2 = doc.add_table(rows=1, cols=1); no_b(t2)
    cell_bg(t2.cell(0,0), bg)
    p2 = t2.cell(0,0).paragraphs[0]
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after  = Pt(10)
    p2.paragraph_format.left_indent  = Cm(0.5)
    p2.paragraph_format.right_indent = Cm(0.3)
    sf(p2.add_run(corps), 11, False, DARK)
    sp(8)

def qa(question, reponse):
    """Q&A : question fond doré, réponse fond vert clair"""
    t = doc.add_table(rows=1, cols=1); no_b(t)
    cell_bg(t.cell(0,0), "7A4800")
    p = t.cell(0,0).paragraphs[0]
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after  = Pt(7)
    p.paragraph_format.left_indent  = Cm(0.4)
    sf(p.add_run("❓  " + question), 11, True, WHITE)
    t2 = doc.add_table(rows=1, cols=1); no_b(t2)
    cell_bg(t2.cell(0,0), LGREEN)
    p2 = t2.cell(0,0).paragraphs[0]
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after  = Pt(10)
    p2.paragraph_format.left_indent  = Cm(0.5)
    p2.paragraph_format.right_indent = Cm(0.3)
    sf(p2.add_run("✅  "), 11, True, GREEN)
    sf(p2.add_run(reponse), 11, False, DARK)
    sp(8)

def code_block(code_txt):
    t = doc.add_table(rows=1, cols=1); no_b(t)
    cell_bg(t.cell(0,0), "1E1E2E")
    p = t.cell(0,0).paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(8)
    p.paragraph_format.left_indent  = Cm(0.5)
    r = p.add_run(code_txt)
    r.font.name = "Courier New"; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xA8,0xFF,0xC2)
    sp(8)

def analogy(txt):
    t = doc.add_table(rows=1, cols=1); no_b(t)
    cell_bg(t.cell(0,0), LYELL)
    p = t.cell(0,0).paragraphs[0]
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after  = Pt(7)
    p.paragraph_format.left_indent  = Cm(0.4)
    sf(p.add_run("💡  Analogie : "), 10, True, GOLD)
    sf(p.add_run(txt), 10, False, DARK, True)
    sp(8)

# ══════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════
t_hdr = doc.add_table(rows=1, cols=1); no_b(t_hdr)
cell_bg(t_hdr.cell(0,0), NAVY_HEX)
ph = t_hdr.cell(0,0).paragraphs[0]
ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
ph.paragraph_format.space_before = Pt(16); ph.paragraph_format.space_after = Pt(6)
sf(ph.add_run("COMPRENDRE DJANGO REST FRAMEWORK"), 20, True, WHITE)
ph2 = t_hdr.cell(0,0).add_paragraph()
ph2.alignment = WD_ALIGN_PARAGRAPH.CENTER
ph2.paragraph_format.space_after = Pt(16)
sf(ph2.add_run("Guide de préparation  •  Soutenance GABIAM Dédé Honorine"), 10, False, RGBColor(0xD4,0xA0,0x17), True)

t_gold = doc.add_table(rows=1, cols=1); no_b(t_gold)
cell_bg(t_gold.cell(0,0), GOLD_HEX)
tg = t_gold.cell(0,0).paragraphs[0]
tg.paragraph_format.space_before = Pt(2); tg.paragraph_format.space_after = Pt(2)
sp(14)

# ══════════════════════════════════════════════════════════════════════
# 1. C'EST QUOI DJANGO ?
# ══════════════════════════════════════════════════════════════════════
section_title("1.  C'est quoi Django ?")

concept("Django en une phrase",
    "Django est un framework Python qui permet de créer rapidement un backend "
    "(serveur, base de données, API) sans repartir de zéro à chaque fois.")

analogy("Django c'est comme une cuisine déjà équipée. Tu n'as pas besoin de "
        "fabriquer les couteaux ou le four — tout est là, tu cuisines directement.")

concept("Django REST Framework (DRF) — ce qu'on a utilisé",
    "DRF est une extension de Django qui permet de créer des API REST. "
    "Une API REST c'est un ensemble d'URLs (endpoints) qui retournent des données JSON "
    "plutôt que des pages HTML. C'est ce que notre frontend React consomme.")

analogy("L'API Django c'est comme un serveur de restaurant. "
        "React (le frontend) est le client qui commande. "
        "Django prend la commande, va chercher les données en base, et retourne la réponse.")

# ══════════════════════════════════════════════════════════════════════
# 2. LES CONCEPTS CLÉS
# ══════════════════════════════════════════════════════════════════════
section_title("2.  Les 4 concepts clés de Django REST")

concept("① Le Model — la structure des données en base",
    "Un Model Django correspond à une table en base de données. "
    "Chaque attribut du model = une colonne dans la table. "
    "Par exemple, le model StakingPosition a les champs : user, token_address, amount, "
    "duration_years, end_date, tx_hash. Django crée automatiquement la table SQL correspondante.",
    LBLUE)

code_block(
    "class StakingPosition(models.Model):\n"
    "    user           = models.ForeignKey(User, on_delete=models.CASCADE)\n"
    "    token_address  = models.CharField(max_length=42)\n"
    "    amount         = models.DecimalField(...)\n"
    "    duration_years = models.IntegerField()\n"
    "    end_date       = models.DateField()\n"
    "    tx_hash        = models.CharField(max_length=66)"
)

concept("② Le Serializer — la conversion entre Python et JSON",
    "Le serializer transforme un objet Python (Model) en JSON pour l'envoyer au frontend, "
    "et transforme le JSON reçu du frontend en objet Python pour le sauvegarder. "
    "Il valide aussi les données (champ obligatoire, format correct, etc.).",
    LBLUE)

analogy("Le serializer c'est comme un traducteur entre Python et JSON. "
        "Sans lui, Django ne saurait pas comment présenter les données au frontend.")

concept("③ La Vue (ViewSet) — la logique métier",
    "La vue reçoit la requête HTTP (GET, POST, PUT, DELETE), applique la logique "
    "(vérifier que l'utilisateur est connecté, calculer end_date, sauvegarder...) "
    "et retourne la réponse JSON via le serializer.",
    LBLUE)

code_block(
    "# Quand le frontend fait POST /api/blockchain/stake-tokens/stake/\n"
    "# Django exécute cette fonction :\n"
    "def stake(self, request):\n"
    "    serializer = StakeSerializer(data=request.data)\n"
    "    if serializer.is_valid():\n"
    "        serializer.save(user=request.user)  # sauvegarde en base\n"
    "        return Response(serializer.data, status=201)\n"
    "    return Response(serializer.errors, status=400)"
)

concept("④ L'URL (Router) — le routage des requêtes",
    "Le router associe chaque URL à une vue. "
    "Par exemple : POST /api/blockchain/stake-tokens/stake/ → vue StakingViewSet.stake(). "
    "C'est comme un standard téléphonique qui redirige l'appel vers le bon interlocuteur.",
    LBLUE)

# ══════════════════════════════════════════════════════════════════════
# 3. JWT — AUTHENTIFICATION
# ══════════════════════════════════════════════════════════════════════
section_title("3.  L'authentification JWT")

concept("Pourquoi JWT ?",
    "HTTP est sans état (stateless) : le serveur ne se souvient pas de toi entre deux requêtes. "
    "JWT permet de prouver ton identité à chaque requête sans stocker de session côté serveur. "
    "Tu envoies un token signé dans chaque requête, Django le vérifie et sait qui tu es.")

concept("Access token vs Refresh token",
    "Access token : durée courte (~5 min). Envoyé dans chaque requête API. "
    "Si quelqu'un le vole, il expire vite.\n\n"
    "Refresh token : durée longue (~7 jours). Utilisé uniquement pour obtenir "
    "un nouvel access token quand celui-ci expire. Jamais envoyé aux endpoints normaux.")

analogy("L'access token c'est un badge journalier qui expire le soir. "
        "Le refresh token c'est ta carte d'employé qui te permet d'obtenir un nouveau badge le lendemain.")

concept("Comment on l'utilise dans notre code (blockchainApi.ts)",
    "Chaque appel API envoie l'access token dans le header :\n"
    "Authorization: Bearer eyJhbGci...\n\n"
    "Si Django retourne 401 (token expiré), notre apiFetch() récupère automatiquement "
    "le refresh token, appelle /api/accounts/token/refresh/, reçoit un nouvel access token, "
    "et refait la requête — invisible pour l'utilisateur.")

code_block(
    "# Dans blockchainApi.ts — gestion automatique du refresh\n"
    "if (res.status === 401) {\n"
    "    const refreshRes = await fetch('/api/accounts/token/refresh/', {\n"
    "        method: 'POST',\n"
    "        body: JSON.stringify({ refresh: getRefreshToken() })\n"
    "    });\n"
    "    const data = await refreshRes.json();\n"
    "    setTokens(data.access, refresh);  // nouveau token stocké\n"
    "    // on refait la requête originale\n"
    "}"
)

# ══════════════════════════════════════════════════════════════════════
# 4. NOS ENDPOINTS — CONCRETS
# ══════════════════════════════════════════════════════════════════════
section_title("4.  Les endpoints Django qu'on a utilisés")

concept("Tableau des principaux endpoints",
    "GET  /api/blockchain/stake-tokens/          → mes positions de staking\n"
    "POST /api/blockchain/stake-tokens/stake/    → enregistrer un nouveau stake\n"
    "POST /api/blockchain/unstake/               → enregistrer un unstake\n"
    "GET  /api/blockchain/networks/              → liste des réseaux (Sepolia, etc.)\n"
    "GET  /api/blockchain/token-prices/          → prix des tokens en USD\n"
    "POST /api/accounts/token/                   → login → retourne access + refresh token\n"
    "POST /api/accounts/token/refresh/           → renouveler l'access token\n"
    "POST /api/payments/checkout/                → créer un checkout FedaPay\n"
    "GET  /api/academy/courses/                  → liste des cours publiés")

# ══════════════════════════════════════════════════════════════════════
# 5. CYCLE D'UNE REQUÊTE — EXEMPLE CONCRET
# ══════════════════════════════════════════════════════════════════════
section_title("5.  Cycle complet d'une requête — exemple : enregistrer un stake")

concept("Ce qui se passe quand on appelle stakingApi.recordStake()",
    "① React envoie POST /api/blockchain/stake-tokens/stake/ avec le JSON :\n"
    "   { token_address, token_symbol, chain_id, amount, tx_hash, duration_years }\n\n"
    "② Django reçoit la requête, vérifie le JWT dans le header Authorization\n\n"
    "③ Le Router Django redirige vers la vue StakingViewSet.stake()\n\n"
    "④ Le Serializer valide les données (champs présents ? formats corrects ?)\n\n"
    "⑤ La vue calcule end_date = aujourd'hui + duration_years\n\n"
    "⑥ Django sauvegarde la StakingPosition en base PostgreSQL\n\n"
    "⑦ Django retourne 201 Created + l'objet créé en JSON\n\n"
    "⑧ React reçoit la réponse et rafraîchit l'affichage")

# ══════════════════════════════════════════════════════════════════════
# 6. ERREURS COURANTES ET LEUR SIGNIFICATION
# ══════════════════════════════════════════════════════════════════════
section_title("6.  Les erreurs qu'on a rencontrées et ce qu'elles veulent dire")

concept("400 Bad Request",
    "Les données envoyées sont invalides selon le serializer. "
    "Exemple : champ manquant, mauvais format, valeur inconnue. "
    "Django retourne un JSON avec le détail de l'erreur : { 'chain_id': ['Network not found'] }",
    "FFF0F0")

concept("401 Unauthorized",
    "Le JWT est absent, invalide ou expiré. "
    "Notre apiFetch() intercepte ce cas et tente un refresh automatique. "
    "Si le refresh échoue aussi, l'utilisateur doit se reconnecter.",
    "FFF0F0")

concept("403 Forbidden",
    "L'utilisateur est authentifié mais n'a pas la permission. "
    "Exemple : un utilisateur normal qui essaie d'accéder à un endpoint admin-only.",
    "FFF0F0")

concept("502 Bad Gateway",
    "Le serveur Django est planté ou pas démarré. "
    "En développement : oublier de lancer python manage.py runserver. "
    "En production : le service Render ou Heroku est en veille.",
    "FFF0F0")

# ══════════════════════════════════════════════════════════════════════
# 7. QUESTIONS / RÉPONSES
# ══════════════════════════════════════════════════════════════════════
section_title("7.  Questions probables à la soutenance")

qa("C'est quoi Django et pourquoi l'avoir utilisé ?",
   "Django est un framework Python pour créer des backends et des API. "
   "On l'a utilisé parce que le backend était déjà construit avec Django par le développeur backend. "
   "Mon rôle côté frontend était de consommer cette API : envoyer des requêtes HTTP et "
   "traiter les réponses JSON. J'ai dû apprendre Django REST Framework pour comprendre "
   "comment les données étaient structurées et pourquoi certaines requêtes échouaient.")

qa("Comment Django sait que c'est bien toi qui fait la requête ?",
   "Via le JWT. À la connexion, Django génère deux tokens signés avec une clé secrète. "
   "À chaque requête, le frontend envoie l'access token dans le header Authorization. "
   "Django le décode, vérifie la signature et extrait l'identité de l'utilisateur. "
   "Si le token est modifié ou expiré, Django retourne 401.")

qa("Quelle est la différence entre GET et POST ?",
   "GET sert à lire des données (ex: liste des cours, mes positions de staking). "
   "Il n'y a pas de corps dans la requête. "
   "POST sert à créer ou envoyer des données (ex: enregistrer un stake, créer un checkout). "
   "Les données sont envoyées dans le corps de la requête en JSON.")

qa("C'est quoi un serializer Django et pourquoi c'est important ?",
   "Le serializer fait deux choses : il convertit les objets Python en JSON pour les envoyer "
   "au frontend (sérialisation), et il valide et convertit le JSON reçu du frontend "
   "en objets Python pour les sauvegarder (désérialisation). "
   "Sans serializer, on devrait écrire manuellement la conversion et la validation pour chaque champ.")

qa("Comment avez-vous appris Django pendant le stage ?",
   "Je n'avais pas de connaissances Django au départ. Avec l'accord de mon tuteur M. Hospice KAKE, "
   "j'ai suivi des formations sur Udemy sur Django REST Framework et l'authentification JWT. "
   "J'ai appris en pratique en déboguant les erreurs réelles : lire les messages d'erreur 400, "
   "comprendre les serializers, et dialoguer précisément avec le développeur backend "
   "pour ajuster les structures de données.")

qa("Quelle différence entre votre base de données et la blockchain ?",
   "La base de données Django (PostgreSQL) stocke des données applicatives : "
   "qui a staké quoi, depuis quand, quelle durée. C'est modifiable et centralisé — "
   "c'est nous qui la contrôlons. "
   "La blockchain stocke les vraies transactions financières : elle est immuable, "
   "décentralisée et publique. Les deux sont complémentaires : "
   "la blockchain garantit la transaction, Django garde l'historique enrichi.")

qa("Qu'est-ce que CORS et pourquoi ça a pu poser problème ?",
   "CORS (Cross-Origin Resource Sharing) est une règle de sécurité du navigateur : "
   "il refuse qu'un site (stake-easy.vercel.app) appelle une API sur un autre domaine "
   "(golden-backend.onrender.com) sans autorisation explicite. "
   "Le développeur Django a dû configurer django-cors-headers pour autoriser "
   "les requêtes venant de notre frontend.")

# ══════════════════════════════════════════════════════════════════════
# RÉSUMÉ FINAL
# ══════════════════════════════════════════════════════════════════════
section_title("8.  Résumé — ce qu'il faut absolument retenir")

t_sum = doc.add_table(rows=5, cols=2)
no_b(t_sum)
rows_data = [
    ("Django",         "Framework Python pour créer des API backend"),
    ("DRF",            "Extension Django pour faire des API REST (JSON)"),
    ("Model",          "Structure d'une table en base de données"),
    ("Serializer",     "Traducteur Python ↔ JSON + validation des données"),
    ("JWT",            "Système de tokens pour s'authentifier sans session"),
]
for i, (k, v) in enumerate(rows_data):
    bg = "E8F0FB" if i % 2 == 0 else "F5F8FF"
    cell_bg(t_sum.rows[i].cells[0], BLUE_HEX)
    cell_bg(t_sum.rows[i].cells[1], bg)
    t_sum.rows[i].cells[0].width = Cm(4)
    t_sum.rows[i].cells[1].width = Cm(12)
    pk = t_sum.rows[i].cells[0].paragraphs[0]
    pk.paragraph_format.space_before = Pt(6); pk.paragraph_format.space_after = Pt(6)
    pk.paragraph_format.left_indent = Cm(0.3)
    sf(pk.add_run(k), 10, True, WHITE)
    pv = t_sum.rows[i].cells[1].paragraphs[0]
    pv.paragraph_format.space_before = Pt(6); pv.paragraph_format.space_after = Pt(6)
    pv.paragraph_format.left_indent = Cm(0.3)
    sf(pv.add_run(v), 10, False, DARK)

sp(14)
t_ft = doc.add_table(rows=1, cols=1); no_b(t_ft)
cell_bg(t_ft.cell(0,0), NAVY_HEX)
pf = t_ft.cell(0,0).paragraphs[0]
pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
pf.paragraph_format.space_before = Pt(10); pf.paragraph_format.space_after = Pt(10)
sf(pf.add_run("GABIAM Dédé Honorine  •  StakeEasy / GoldenBridge  •  Soutenance 2025-2026"), 9, False, WHITE, True)

doc.save(OUT)
print(f"✅ {OUT}")
